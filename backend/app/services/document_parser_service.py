import io
import re
import httpx
import pypdf
import docx
import openpyxl


def extract_text(file_bytes: bytes, filename: str, content_type: str | None = None) -> str:
    """Extract raw text from various document formats to send to OpenAI for analysis."""
    ext = filename.lower()
    
    # 1. Text & CSV Files
    if ext.endswith((".txt", ".csv", ".json", ".xml", ".tsv", ".yaml", ".yml", ".md")):
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1")
            
    # 2. PDF Documents
    elif ext.endswith(".pdf") or content_type == "application/pdf":
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            text = []
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text.append(t)
            return "\n".join(text)
        except Exception as e:
            return f"[Error parsing PDF: {e}]"
            
    # 3. Word Documents
    elif ext.endswith(".docx") or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            return "\n".join([p.text for p in doc.paragraphs])
        except Exception as e:
            return f"[Error parsing Word document: {e}]"
            
    # 4. Excel Spreadsheets
    elif ext.endswith((".xlsx", ".xls")) or content_type in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel"
    ):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            text = []
            for sheet in wb.worksheets:
                text.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(v) if v is not None else "" for v in row]
                    if any(row_vals):
                        text.append(" | ".join(row_vals))
            return "\n".join(text)
        except Exception as e:
            return f"[Error parsing Excel file: {e}]"
            
    # 5. Fallback - decode as string anyway
    try:
        return file_bytes.decode("utf-8", errors="ignore")
    except Exception:
        return ""


def extract_url_text(url: str) -> str:
    """Fetch an online link (e.g. Google Sheets/Docs) and extract visible text from the page."""
    try:
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        }
        resp = httpx.get(url, headers=headers, timeout=15, follow_redirects=True)
        if resp.status_code != 200:
            return f"[Failed to fetch link: HTTP {resp.status_code}]"
            
        html = resp.text
        # Strip script & style tags
        html = re.sub(r"<(script|style).*?>.*?</\1>", "", html, flags=re.DOTALL | re.IGNORECASE)
        # Strip all HTML tags
        text = re.sub(r"<.*?>", " ", html)
        # Clean up whitespace
        text = re.sub(r"\s+", " ", text).strip()
        # Cap text size to prevent prompt overflow
        return text[:15000]
    except Exception as e:
        return f"[Error fetching URL text: {e}]"
